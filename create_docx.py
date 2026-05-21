from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

document = Document()

# Title
title = document.add_heading('SmartLab Monitoring System', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Aim
document.add_heading('Aim', level=1)
document.add_paragraph('To ensure academic integrity during lab sessions and exams by establishing a transparent, real-time monitoring environment that tracks application usage and provides immediate feedback on unauthorized activities.')

# Objectives
document.add_heading('Objectives', level=1)
objectives = [
    'Real-Time Monitoring: Continuously track the active foreground windows on student machines using a lightweight desktop client.',
    'Automated Rule Enforcement: Dynamically evaluate student activity against permitted subjects and allowed applications.',
    'Transparent Notifications: Instantly alert students via persistent UI banners and toast notifications when an unauthorized window is detected.',
    'Centralized Administration: Provide teachers with a dashboard to oversee ongoing sessions, manage dynamic subjects, and review comprehensive student activity logs.'
]
for obj in objectives:
    document.add_paragraph(obj, style='List Bullet')

# Method & Flow Diagram
document.add_heading('Method & Flow Diagram', level=1)
document.add_paragraph('The operational architecture of the SmartLab Monitoring System is built upon a continuous, low-latency feedback loop between a specialized desktop client and a centralized backend server.')

document.add_heading('Core Operational Methodology', level=2)
method_steps = [
    'Continuous Data Collection: The student operates a desktop application (built with Electron and Node.js). It periodically captures the title and metadata of the currently active foreground window on the operating system.',
    'Real-Time Transmission: Using a persistent WebSocket connection, the desktop client securely transmits this active window data, along with a unique student identifier, to the Node.js backend.',
    'Dynamic Rule Evaluation: Upon receiving the data, the backend server processes it through its rules engine. The server dynamically fetches the allowed applications and URLs based on the subject the student is currently enrolled in.',
    'Immediate Action & Notification: \n - Authorized Activity: If the window is deemed acceptable, the server logs a "Normal" state. The teacher\'s dashboard updates silently.\n - Unauthorized Activity: If an unapproved application is detected, the server instantly triggers a violation. It sends a WebSocket alert back to the student\'s app, which renders a persistent warning banner and a toast notification explaining the violation. Simultaneously, the event is recorded in the permanent audit log and prominently displayed on the teacher\'s dashboard.'
]
for step in method_steps:
    document.add_paragraph(step, style='List Number')

document.add_paragraph('\nFlow Diagram:')
flow_diagram_text = """1. Student Desktop App -> Captures Active Foreground Window
2. Activity Payload -> Sent via WebSocket to Backend Node.js Server
3. Rules Engine -> Evaluates if the Window is Authorized for the Current Subject
4. If Not Authorized -> Triggers Violation Sequence (Sends Socket Alert to Student UI, Adds Violation to Audit Log)
5. If Authorized -> Records Normal Activity
6. Audit Logs & Normal Activity -> Updates Teacher Admin Dashboard
7. Teacher Admin Dashboard -> Provides Live Oversight & Proctor Action"""
document.add_paragraph(flow_diagram_text)

# Results & Deliverables
document.add_heading('Results & Deliverables', level=1)
results = [
    'Student Desktop Client: An application capable of discreet system monitoring and displaying immediate violation feedback.',
    'Backend Infrastructure: A robust Express.js/Socket.io server handling real-time connections, session state, and dynamic subject configurations.',
    'Teacher Dashboard: A web interface displaying active sessions, live system logs, and security audits for proctors.',
    'Notification System: A fully integrated alert system ensuring students know exactly why and when they were flagged.'
]
for res in results:
    document.add_paragraph(res, style='List Bullet')

# Conclusion
document.add_heading('Conclusion', level=1)
document.add_paragraph('The SmartLab Monitoring System successfully mitigates academic dishonesty by combining real-time surveillance with immediate, transparent feedback. Rather than acting purely as a punitive tool, its instant notifications correct behavior on the spot, creating a fairer, more secure, and easily manageable proctored lab environment for educators and students alike.')

document.save('SmartLab_Project_Summary_Expanded.docx')
print("Successfully updated SmartLab_Project_Summary_Expanded.docx")
